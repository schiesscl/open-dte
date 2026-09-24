from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
import openpyxl
import os
from inventario.models import Producto


def safe_int(value):
    """Convierte un valor a entero de forma segura, retorna 0 si falla"""
    if not value:
        return 0
    try:
        if isinstance(value, str):
            value = value.strip()
        return int(float(value)) if value else 0
    except (ValueError, TypeError):
        return 0


class Command(BaseCommand):
    help = 'Importa productos desde un archivo Excel'

    def add_arguments(self, parser):
        parser.add_argument(
            'archivo',
            type=str,
            help='Ruta del archivo Excel a importar'
        )
        parser.add_argument(
            '--limpiar',
            action='store_true',
            help='Elimina todos los productos antes de importar'
        )

    def handle(self, *args, **options):
        archivo = options['archivo']
        limpiar = options['limpiar']

        # Verificar que el archivo existe
        if not os.path.exists(archivo):
            raise CommandError(f'El archivo {archivo} no existe')

        # Limpiar productos si se solicita
        if limpiar:
            Producto.objects.all().delete()
            self.stdout.write(self.style.SUCCESS('[OK] Productos anteriores eliminados'))

        try:
            # Cargar workbook
            wb = openpyxl.load_workbook(archivo, data_only=True)
            ws = wb.active

            total_rows = ws.max_row - 1  # Menos encabezado

            self.stdout.write(f'\nImportando {total_rows} productos...\n')

            # Procesar productos
            with transaction.atomic():
                productos_creados = 0
                productos_actualizados = 0
                errores = 0

                for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
                    try:
                        # Mapeo de columnas: [0]=SUCURSAL (ignorado), [1]=BODEGA (ignorado),
                        # [2]=COD.PRODUCTO, [3]=COD.ALTERNATIVO, [4]=DESCRIPCION,
                        # [5]=CANTIDAD, [6]=REAL, [7]=DIFERENCIA, [8]=MINIMO REPOSICION
                        codigo = row[2]
                        codigo_alt = row[3] or ''
                        descripcion = row[4]
                        cantidad = safe_int(row[5])
                        stock_real = safe_int(row[6])
                        stock_minimo = safe_int(row[8])

                        if not codigo or not descripcion:
                            continue

                        # Usar update_or_create para crear o actualizar
                        producto, created = Producto.objects.update_or_create(
                            codigo=codigo,
                            defaults={
                                'codigo_alternativo': codigo_alt,
                                'descripcion': descripcion,
                                'stock_actual': cantidad,
                                'stock_real': stock_real,
                                'stock_minimo': stock_minimo,
                                'activo': True,
                            }
                        )

                        if created:
                            productos_creados += 1
                        else:
                            productos_actualizados += 1

                        # Mostrar progreso cada 50 productos
                        if (productos_creados + productos_actualizados) % 50 == 0:
                            self.stdout.write(f'  Procesados: {productos_creados + productos_actualizados}/{total_rows}')

                    except Exception as e:
                        errores += 1
                        self.stdout.write(
                            self.style.ERROR(f'[ERROR] Error en fila {row_idx}: {str(e)[:100]}')
                        )

            # Resumen
            self.stdout.write('\n' + '='*60)
            self.stdout.write(self.style.SUCCESS(f'[OK] IMPORTACIÓN COMPLETADA'))
            self.stdout.write(f'  Productos creados:    {productos_creados}')
            self.stdout.write(f'  Productos actualizados: {productos_actualizados}')
            self.stdout.write(f'  Total: {Producto.objects.count()} productos en BD')
            if errores > 0:
                self.stdout.write(self.style.WARNING(f'  [WARNING] Errores: {errores}'))
            self.stdout.write('='*60 + '\n')

        except Exception as e:
            raise CommandError(f'Error al procesar Excel: {str(e)}')
